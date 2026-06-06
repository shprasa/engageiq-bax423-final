"""Update brief.docx with current dataset and benchmark stats."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document

CODE_DIR = Path(__file__).resolve().parent.parent
PROJECT_ROOT = CODE_DIR.parent
BRIEF = PROJECT_ROOT / "brief.docx"
BENCHMARK = PROJECT_ROOT / "data" / "benchmark_results.json"


def _replace_in_paragraph(paragraph, old: str, new: str) -> bool:
    if old not in paragraph.text:
        return False
    text = paragraph.text.replace(old, new)
    paragraph.clear()
    paragraph.add_run(text)
    return True


def main() -> None:
    if not BRIEF.exists():
        raise SystemExit(f"Missing {BRIEF}")

    stats = json.loads(BENCHMARK.read_text(encoding="utf-8"))
    ds = stats["dataset"]
    lb = stats["learning_benchmark"]
    rows = ds["dataset_rows"]
    gh = ds["sources"]["github"]
    gha = ds["sources"]["gharchive"]
    rows_fmt = f"{rows:,}"
    rl_without = lb["avg_reward_last10_without_rl"]
    rl_with = lb["avg_reward_last10_with_rl"]
    rl_gain = lb["reward_improvement_last10"]
    cum_with = lb["cumulative_reward_with_rl"]
    cum_without = lb["cumulative_reward_without_rl"]

    doc = Document(str(BRIEF))

    for p in doc.paragraphs:
        _replace_in_paragraph(
            p,
            "Snapshot: 12,750 records (100% live API URLs), all 15 domains, in data/opportunities_snapshot.csv.",
            f"Snapshot: {rows_fmt} records (100% live API URLs), all 15 domains, in data/opportunities_snapshot.csv "
            f"({gh:,} GitHub API + {gha:,} GitHub Archive).",
        )
        _replace_in_paragraph(
            p,
            "Over {rounds} simulated rounds, average reward in the last 10 improves from 0.07 to 1.00 (+0.93),",
            f"Over {lb['rounds']} simulated rounds, average reward in the last 10 improves from "
            f"{rl_without:.2f} to {rl_with:.2f} (+{rl_gain:.2f}),",
        )

    # Table 0 — source mix
    t0 = doc.tables[0]
    t0.rows[1].cells[1].text = f"{gh:,}"
    t0.rows[1].cells[2].text = f"{gh:,}"
    t0.rows[2].cells[1].text = f"{gha:,}"
    t0.rows[2].cells[2].text = f"{gha:,}"
    t0.rows[3].cells[1].text = rows_fmt
    t0.rows[3].cells[2].text = rows_fmt

    # Table 4 — RL benchmark
    t4 = doc.tables[4]
    t4.rows[1].cells[1].text = f"{rl_with:.2f}"
    t4.rows[1].cells[2].text = f"{rl_without:.2f}"
    t4.rows[1].cells[3].text = f"+{rl_gain:.2f}"
    t4.rows[2].cells[1].text = f"{cum_with:.0f}"
    t4.rows[2].cells[2].text = f"{cum_without:.2f}".rstrip("0").rstrip(".")

    doc.save(str(BRIEF))
    print(f"Updated {BRIEF}")


if __name__ == "__main__":
    main()
